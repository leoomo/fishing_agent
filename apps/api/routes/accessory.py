"""
钓鱼配件管理 API 路由

提供钓鱼配件的 CRUD 操作和初始化数据功能
包括钩子、铅坠、转环、前导线、浮漂、别针等
"""

import logging
from io import BytesIO
from typing import List, Optional

from fastapi import APIRouter, Body, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import StreamingResponse
from sqlalchemy import func

from apps.api.auth.dependencies import get_current_user, require_permission
from apps.api.models import Accessory
from apps.api.orm import get_db_session
from apps.api.schemas.accessory import (
    AccessoryBatchDeleteResponse,
    AccessoryCategoryEnum,
    AccessoryCategoryStats,
    AccessoryCategoryStatsResponse,
    AccessoryCreate,
    AccessoryImportPreviewResponse,
    AccessoryImportPreviewItem,
    AccessoryImportResult,
    AccessoryInitDataResponse,
    AccessoryListItem,
    AccessoryListResponse,
    AccessoryOptionsResponse,
    AccessoryResponse,
    AccessoryUpdate,
    UserLevelEnum,
)

logger = logging.getLogger(__name__)

router = APIRouter()

# 分类配置
CATEGORY_CONFIG = {
    "hook": {"label": "钩子", "icon": "🪝", "color": "#1890ff"},
    "sinker": {"label": "铅坠", "icon": "⚓", "color": "#722ed1"},
    "swivel": {"label": "转环", "icon": "🔗", "color": "#52c41a"},
    "leader": {"label": "前导线", "icon": "〰️", "color": "#faad14"},
    "float": {"label": "浮漂", "icon": "🔴", "color": "#eb2f96"},
    "snap": {"label": "别针", "icon": "📎", "color": "#13c2c2"},
    "other": {"label": "其他", "icon": "🎣", "color": "#8c8c8c"},
}

# 用户等级配置
USER_LEVEL_CONFIG = {
    "beginner": {"label": "新手", "color": "#52c41a"},
    "intermediate": {"label": "进阶", "color": "#1890ff"},
    "advanced": {"label": "高级", "color": "#722ed1"},
}

# 常用材质
COMMON_MATERIALS = [
    "碳钢",
    "不锈钢",
    "钨合金",
    "铅",
    "尼龙",
    "碳素",
    "氟碳",
    "钛合金",
    "黄铜",
    "PE材质",
]

# 初始化数据
DEFAULT_ACCESSORIES = [
    # 钩子类
    {
        "name": "曲柄钩",
        "category": "hook",
        "description": "路亚软饵专用钩，钩柄弯曲设计便于挂装软饵",
        "features": "防挂底、软饵专用、多种号数",
        "size": "#1/0 - #5/0",
        "material": "碳钢",
        "target_species": "黑鲈、鳜鱼",
        "applicable_rigs": "Texas钓组、Carolina钓组",
        "user_level": "beginner",
    },
    {
        "name": "铅头钩",
        "category": "hook",
        "description": "带铅头的钩子，可直接挂软饵使用",
        "features": "自带配重、快速下沉、操作简单",
        "size": "3.5g - 21g",
        "weight": 7.0,
        "material": "铅+碳钢",
        "target_species": "黑鲈、鳜鱼、翘嘴",
        "applicable_rigs": "铅头钩钓组",
        "user_level": "beginner",
    },
    {
        "name": "虫钩",
        "category": "hook",
        "description": "适合挂装细长型软饵的钩子",
        "features": "细长钩条、适合面条虫",
        "size": "#1 - #4",
        "material": "碳钢",
        "target_species": "黑鲈、鲤鱼",
        "applicable_rigs": "Wacky钓组、Neko钓组",
        "user_level": "intermediate",
    },
    {
        "name": "三本钩",
        "category": "hook",
        "description": "硬饵常用的三叉钩",
        "features": "三个钩尖、中鱼率高、适合硬饵",
        "size": "#4 - #1/0",
        "material": "碳钢",
        "target_species": "各类路亚对象鱼",
        "applicable_rigs": "硬饵钓组",
        "user_level": "beginner",
    },
    # 铅坠类
    {
        "name": "子弹铅",
        "category": "sinker",
        "description": "子弹形状的铅坠，可穿过障碍区",
        "features": "子弹造型、防挂底、穿透力强",
        "size": "3.5g - 21g",
        "weight": 7.0,
        "material": "铅",
        "target_species": "黑鲈、鳜鱼",
        "applicable_rigs": "Texas钓组",
        "user_level": "beginner",
    },
    {
        "name": "钨钢子弹铅",
        "category": "sinker",
        "description": "钨合金材质的子弹铅，体积更小密度更大",
        "features": "高密度、体积小、灵敏度高、环保",
        "size": "3.5g - 14g",
        "weight": 7.0,
        "material": "钨合金",
        "target_species": "黑鲈、鳜鱼",
        "applicable_rigs": "Texas钓组、Carolina钓组",
        "user_level": "intermediate",
        "price_min": 15,
        "price_max": 50,
    },
    {
        "name": "Drop Shot铅",
        "category": "sinker",
        "description": "倒吊钓组专用的圆柱形铅坠",
        "features": "底部触感好、适合精细作钓",
        "size": "3.5g - 10.5g",
        "weight": 5.0,
        "material": "铅",
        "target_species": "黑鲈、鳜鱼",
        "applicable_rigs": "Drop Shot钓组",
        "user_level": "intermediate",
    },
    {
        "name": "夹铅",
        "category": "sinker",
        "description": "可夹在线上的小铅粒",
        "features": "可调节位置、微调配重",
        "size": "B - 4B",
        "weight": 0.5,
        "material": "铅",
        "target_species": "各类对象鱼",
        "applicable_rigs": "多种钓组",
        "user_level": "beginner",
    },
    # 转环类
    {
        "name": "八字环",
        "category": "swivel",
        "description": "连接主线和前导线的基础转环",
        "features": "防缠绕、连接稳固",
        "size": "#3 - #7",
        "material": "不锈钢",
        "target_species": "各类对象鱼",
        "applicable_rigs": "多种钓组",
        "user_level": "beginner",
    },
    {
        "name": "滚珠轴承转环",
        "category": "swivel",
        "description": "内置滚珠轴承的高端转环",
        "features": "超顺滑旋转、适合亮片等旋转饵",
        "size": "#1 - #5",
        "material": "不锈钢+黄铜",
        "target_species": "翘嘴、鲈鱼",
        "applicable_rigs": "亮片钓组",
        "user_level": "intermediate",
        "price_min": 5,
        "price_max": 20,
    },
    # 前导线类
    {
        "name": "碳素前导线",
        "category": "leader",
        "description": "氟碳材质的前导线，水下近乎隐形",
        "features": "高透明、耐磨、抗紫外线",
        "size": "6lb - 20lb",
        "material": "氟碳",
        "target_species": "各类对象鱼",
        "applicable_rigs": "多种钓组",
        "user_level": "beginner",
    },
    {
        "name": "钢丝前导线",
        "category": "leader",
        "description": "防咬断的钢丝前导线",
        "features": "防咬断、适合有牙齿的鱼种",
        "size": "20lb - 80lb",
        "material": "不锈钢丝",
        "target_species": "狗鱼、鳡鱼、带鱼",
        "applicable_rigs": "多种钓组",
        "user_level": "intermediate",
    },
    # 浮漂类
    {
        "name": "路亚浮漂",
        "category": "float",
        "description": "路亚专用浮漂，用于控制饵的泳层",
        "features": "控制泳层、增加抛投距离",
        "size": "3g - 15g",
        "weight": 8.0,
        "material": "泡沫+塑料",
        "target_species": "翘嘴、鲈鱼",
        "applicable_rigs": "浮漂钓组",
        "user_level": "intermediate",
    },
    {
        "name": "水滴浮漂",
        "category": "float",
        "description": "水滴形状的透明浮漂",
        "features": "可注水调节重量、隐蔽性好",
        "size": "中号 - 大号",
        "material": "透明塑料",
        "target_species": "翘嘴、鲈鱼",
        "applicable_rigs": "浮漂钓组",
        "user_level": "beginner",
    },
    # 别针类
    {
        "name": "快速别针",
        "category": "snap",
        "description": "快速更换拟饵的连接器",
        "features": "快速更换、操作便捷",
        "size": "#0 - #3",
        "material": "不锈钢",
        "target_species": "各类对象鱼",
        "applicable_rigs": "多种钓组",
        "user_level": "beginner",
    },
    {
        "name": "O型环",
        "category": "snap",
        "description": "连接拟饵的O型分体环",
        "features": "动作自由、适合硬饵",
        "size": "#2 - #4",
        "material": "不锈钢",
        "target_species": "各类对象鱼",
        "applicable_rigs": "硬饵钓组",
        "user_level": "intermediate",
    },
]


# ========== CRUD Endpoints ==========


@router.post(
    "/accessories",
    response_model=AccessoryResponse,
    summary="创建配件",
    status_code=201,
)
async def create_accessory(
    data: AccessoryCreate,
    current_user=Depends(require_permission("content:create")),
):
    """创建新的钓鱼配件"""
    with get_db_session() as session:
        # 检查名称是否已存在
        existing = (
            session.query(Accessory).filter(Accessory.name == data.name).first()
        )
        if existing:
            raise HTTPException(status_code=400, detail=f"配件 '{data.name}' 已存在")

        accessory = Accessory(
            name=data.name,
            category=data.category.value,
            description=data.description,
            features=data.features,
            size=data.size,
            weight=data.weight,
            material=data.material,
            color=data.color,
            quantity_per_pack=data.quantity_per_pack,
            target_species=data.target_species,
            applicable_rigs=data.applicable_rigs,
            best_conditions=data.best_conditions,
            brand=data.brand,
            price_min=data.price_min,
            price_max=data.price_max,
            user_level=data.user_level.value if data.user_level else "beginner",
            image_url=data.image_url,
        )
        session.add(accessory)
        session.commit()
        session.refresh(accessory)

        logger.info(
            f"配件创建成功: accessory_id={accessory.accessory_id}, "
            f"name={accessory.name}, user={current_user.username}"
        )

        return AccessoryResponse.model_validate(accessory)


@router.get(
    "/accessories",
    response_model=AccessoryListResponse,
    summary="获取配件列表",
)
async def list_accessories(
    page: int = Query(1, ge=1, description="页码"),
    page_size: int = Query(20, ge=1, le=100, description="每页数量"),
    category: Optional[str] = Query(None, description="分类筛选"),
    keyword: Optional[str] = Query(None, description="关键词搜索"),
    user_level: Optional[str] = Query(None, description="用户等级筛选"),
    current_user=Depends(get_current_user),
):
    """获取配件分页列表"""
    with get_db_session() as session:
        query = session.query(Accessory)

        # 分类筛选
        if category:
            query = query.filter(Accessory.category == category)

        # 用户等级筛选
        if user_level:
            query = query.filter(Accessory.user_level == user_level)

        # 关键词搜索
        if keyword:
            search_pattern = f"%{keyword}%"
            query = query.filter(
                (Accessory.name.ilike(search_pattern))
                | (Accessory.target_species.ilike(search_pattern))
                | (Accessory.description.ilike(search_pattern))
                | (Accessory.brand.ilike(search_pattern))
            )

        # 统计总数
        total = query.count()

        # 分页
        offset = (page - 1) * page_size
        accessories = (
            query.order_by(Accessory.category, Accessory.name)
            .offset(offset)
            .limit(page_size)
            .all()
        )

        total_pages = (total + page_size - 1) // page_size

        return AccessoryListResponse(
            items=[AccessoryListItem.model_validate(a) for a in accessories],
            total=total,
            page=page,
            page_size=page_size,
            total_pages=total_pages,
        )


@router.get(
    "/accessories/stats",
    response_model=AccessoryCategoryStatsResponse,
    summary="获取分类统计",
)
async def get_category_stats(
    current_user=Depends(get_current_user),
):
    """获取各分类的配件数量统计"""
    with get_db_session() as session:
        stats = (
            session.query(Accessory.category, func.count(Accessory.accessory_id))
            .group_by(Accessory.category)
            .all()
        )

        stats_dict = {category: count for category, count in stats}
        total = sum(stats_dict.values())

        categories = []
        for cat_key, config in CATEGORY_CONFIG.items():
            categories.append(
                AccessoryCategoryStats(
                    category=cat_key,
                    count=stats_dict.get(cat_key, 0),
                    label=config["label"],
                    icon=config["icon"],
                    color=config["color"],
                )
            )

        return AccessoryCategoryStatsResponse(categories=categories, total=total)


@router.get(
    "/accessories/options",
    response_model=AccessoryOptionsResponse,
    summary="获取表单选项",
)
async def get_accessory_options(
    current_user=Depends(get_current_user),
):
    """获取配件表单的选项数据"""
    categories = [
        {"value": key, "label": config["label"], "icon": config["icon"], "color": config["color"]}
        for key, config in CATEGORY_CONFIG.items()
    ]

    user_levels = [
        {"value": key, "label": config["label"], "color": config["color"]}
        for key, config in USER_LEVEL_CONFIG.items()
    ]

    return AccessoryOptionsResponse(
        categories=categories,
        user_levels=user_levels,
        materials=COMMON_MATERIALS,
    )


@router.get(
    "/accessories/{accessory_id}",
    response_model=AccessoryResponse,
    summary="获取配件详情",
)
async def get_accessory(
    accessory_id: int,
    current_user=Depends(get_current_user),
):
    """获取配件详情"""
    with get_db_session() as session:
        accessory = (
            session.query(Accessory)
            .filter(Accessory.accessory_id == accessory_id)
            .first()
        )
        if not accessory:
            raise HTTPException(status_code=404, detail="配件不存在")

        return AccessoryResponse.model_validate(accessory)


@router.put(
    "/accessories/{accessory_id}",
    response_model=AccessoryResponse,
    summary="更新配件",
)
async def update_accessory(
    accessory_id: int,
    data: AccessoryUpdate,
    current_user=Depends(require_permission("content:update")),
):
    """更新配件"""
    with get_db_session() as session:
        accessory = (
            session.query(Accessory)
            .filter(Accessory.accessory_id == accessory_id)
            .first()
        )
        if not accessory:
            raise HTTPException(status_code=404, detail="配件不存在")

        # 检查名称唯一性
        if data.name and data.name != accessory.name:
            existing = (
                session.query(Accessory).filter(Accessory.name == data.name).first()
            )
            if existing:
                raise HTTPException(
                    status_code=400, detail=f"配件 '{data.name}' 已存在"
                )

        # 更新字段
        update_data = data.model_dump(exclude_unset=True)
        if "category" in update_data and update_data["category"]:
            update_data["category"] = update_data["category"].value
        if "user_level" in update_data and update_data["user_level"]:
            update_data["user_level"] = update_data["user_level"].value

        for field, value in update_data.items():
            setattr(accessory, field, value)

        session.commit()
        session.refresh(accessory)

        logger.info(
            f"配件更新成功: accessory_id={accessory_id}, user={current_user.username}"
        )

        return AccessoryResponse.model_validate(accessory)


@router.delete(
    "/accessories/{accessory_id}",
    summary="删除配件",
    status_code=204,
)
async def delete_accessory(
    accessory_id: int,
    current_user=Depends(require_permission("content:delete")),
):
    """删除配件"""
    with get_db_session() as session:
        accessory = (
            session.query(Accessory)
            .filter(Accessory.accessory_id == accessory_id)
            .first()
        )
        if not accessory:
            raise HTTPException(status_code=404, detail="配件不存在")

        name = accessory.name
        session.delete(accessory)
        session.commit()

        logger.info(
            f"配件删除成功: accessory_id={accessory_id}, "
            f"name={name}, user={current_user.username}"
        )


@router.post(
    "/accessories/batch-delete",
    response_model=AccessoryBatchDeleteResponse,
    summary="批量删除配件",
)
async def batch_delete_accessories(
    ids: List[int] = Body(..., embed=True, description="要删除的配件ID列表"),
    current_user=Depends(require_permission("content:delete")),
):
    """批量删除配件"""
    if not ids:
        raise HTTPException(status_code=400, detail="请选择要删除的配件")

    with get_db_session() as session:
        # 查询要删除的记录
        accessories = (
            session.query(Accessory)
            .filter(Accessory.accessory_id.in_(ids))
            .all()
        )

        if not accessories:
            raise HTTPException(status_code=404, detail="未找到要删除的配件")

        deleted_count = len(accessories)
        names = [a.name for a in accessories]

        for accessory in accessories:
            session.delete(accessory)

        session.commit()

        logger.info(
            f"配件批量删除成功: count={deleted_count}, "
            f"names={names}, user={current_user.username}"
        )

        return AccessoryBatchDeleteResponse(
            deleted_count=deleted_count,
            message=f"成功删除 {deleted_count} 个配件",
        )


# ========== Import/Export Endpoints ==========


@router.get(
    "/accessories/export/template",
    summary="下载导入模板",
)
async def download_import_template(
    current_user=Depends(get_current_user),
):
    """下载配件导入Excel模板"""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

        wb = Workbook()
        ws = wb.active
        ws.title = "配件导入模板"

        # 样式定义
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        required_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin'),
        )

        # 定义列
        columns = [
            ("name", "配件名称*", True),
            ("category", "分类*", True),
            ("description", "描述", False),
            ("features", "特点", False),
            ("size", "规格尺寸", False),
            ("weight", "重量(g)", False),
            ("material", "材质", False),
            ("color", "颜色", False),
            ("quantity_per_pack", "每包数量", False),
            ("target_species", "目标鱼种", False),
            ("applicable_rigs", "适用钓组", False),
            ("best_conditions", "最佳条件", False),
            ("brand", "品牌", False),
            ("price_min", "最低价格", False),
            ("price_max", "最高价格", False),
            ("user_level", "用户等级", False),
            ("image_url", "图片URL", False),
        ]

        # 写入表头
        for col_idx, (_, label, required) in enumerate(columns, 1):
            cell = ws.cell(row=1, column=col_idx, value=label)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center')
            cell.border = border
            ws.column_dimensions[cell.column_letter].width = 15

        # 写入说明行
        descriptions = [
            "必填，唯一名称",
            "必填: hook/sinker/swivel/leader/float/snap/other",
            "详细描述",
            "主要特点",
            "如: #1/0, 3.5g",
            "数字",
            "如: 碳钢, 钨合金",
            "如: 银色",
            "数字",
            "如: 黑鲈、鳜鱼",
            "如: Texas钓组",
            "使用条件说明",
            "品牌名称",
            "数字",
            "数字",
            "beginner/intermediate/advanced",
            "图片链接",
        ]
        for col_idx, desc in enumerate(descriptions, 1):
            cell = ws.cell(row=2, column=col_idx, value=desc)
            cell.font = Font(italic=True, color="808080", size=9)
            cell.alignment = Alignment(horizontal='center')
            if columns[col_idx - 1][2]:
                cell.fill = required_fill

        # 写入示例数据
        example_data = [
            "曲柄钩",
            "hook",
            "路亚软饵专用钩",
            "防挂底、软饵专用",
            "#1/0 - #5/0",
            "",
            "碳钢",
            "",
            "10",
            "黑鲈、鳜鱼",
            "Texas钓组、Carolina钓组",
            "",
            "",
            "5",
            "15",
            "beginner",
            "",
        ]
        for col_idx, value in enumerate(example_data, 1):
            cell = ws.cell(row=3, column=col_idx, value=value)
            cell.border = border

        # 保存到内存
        output = BytesIO()
        wb.save(output)
        output.seek(0)

        return StreamingResponse(
            output,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={
                "Content-Disposition": "attachment; filename=accessory_import_template.xlsx"
            },
        )
    except ImportError:
        raise HTTPException(status_code=500, detail="未安装 openpyxl，无法生成模板")


@router.post(
    "/accessories/import/preview",
    response_model=AccessoryImportPreviewResponse,
    summary="预览导入数据",
)
async def preview_import(
    file: UploadFile = File(...),
    current_user=Depends(require_permission("content:create")),
):
    """预览Excel文件中的配件数据"""
    if not file.filename.endswith(('.xlsx', '.xls')):
        raise HTTPException(status_code=400, detail="请上传 Excel 文件 (.xlsx, .xls)")

    try:
        from openpyxl import load_workbook

        content = await file.read()
        wb = load_workbook(BytesIO(content))
        ws = wb.active

        items = []
        valid_count = 0
        invalid_count = 0

        # 获取表头（第一行）
        headers = [cell.value for cell in ws[1]]
        header_map = {
            "配件名称*": "name",
            "分类*": "category",
            "描述": "description",
            "特点": "features",
            "规格尺寸": "size",
            "重量(g)": "weight",
            "材质": "material",
            "颜色": "color",
            "每包数量": "quantity_per_pack",
            "目标鱼种": "target_species",
            "适用钓组": "applicable_rigs",
            "最佳条件": "best_conditions",
            "品牌": "brand",
            "最低价格": "price_min",
            "最高价格": "price_max",
            "用户等级": "user_level",
            "图片URL": "image_url",
        }

        # 跳过表头和说明行，从第3行开始读取数据
        for row_idx, row in enumerate(ws.iter_rows(min_row=3, values_only=True), start=3):
            # 跳过空行
            if not any(row):
                continue

            data = {}
            for col_idx, value in enumerate(row):
                if col_idx < len(headers) and headers[col_idx] in header_map:
                    field = header_map[headers[col_idx]]
                    data[field] = value

            errors = []

            # 验证必填字段
            name = data.get("name")
            category = data.get("category")

            if not name:
                errors.append("配件名称为必填项")
            if not category:
                errors.append("分类为必填项")
            elif category not in ["hook", "sinker", "swivel", "leader", "float", "snap", "other"]:
                errors.append(f"无效的分类: {category}")

            # 验证用户等级
            user_level = data.get("user_level")
            if user_level and user_level not in ["beginner", "intermediate", "advanced"]:
                errors.append(f"无效的用户等级: {user_level}")

            is_valid = len(errors) == 0
            if is_valid:
                valid_count += 1
            else:
                invalid_count += 1

            items.append(AccessoryImportPreviewItem(
                row_number=row_idx,
                name=str(name) if name else "",
                category=str(category) if category else "",
                is_valid=is_valid,
                errors=errors,
                data=data,
            ))

        return AccessoryImportPreviewResponse(
            total_rows=len(items),
            valid_rows=valid_count,
            invalid_rows=invalid_count,
            items=items,
        )
    except ImportError:
        raise HTTPException(status_code=500, detail="未安装 openpyxl，无法处理 Excel 文件")
    except Exception as e:
        logger.error(f"预览导入数据失败: {e}")
        raise HTTPException(status_code=400, detail=f"文件解析失败: {str(e)}")


@router.post(
    "/accessories/import/execute",
    response_model=AccessoryImportResult,
    summary="执行导入",
)
async def execute_import(
    file: UploadFile = File(...),
    current_user=Depends(require_permission("content:create")),
):
    """执行配件数据导入"""
    if not file.filename.endswith(('.xlsx', '.xls')):
        raise HTTPException(status_code=400, detail="请上传 Excel 文件 (.xlsx, .xls)")

    try:
        from openpyxl import load_workbook

        content = await file.read()
        wb = load_workbook(BytesIO(content))
        ws = wb.active

        # 获取表头
        headers = [cell.value for cell in ws[1]]
        header_map = {
            "配件名称*": "name",
            "分类*": "category",
            "描述": "description",
            "特点": "features",
            "规格尺寸": "size",
            "重量(g)": "weight",
            "材质": "material",
            "颜色": "color",
            "每包数量": "quantity_per_pack",
            "目标鱼种": "target_species",
            "适用钓组": "applicable_rigs",
            "最佳条件": "best_conditions",
            "品牌": "brand",
            "最低价格": "price_min",
            "最高价格": "price_max",
            "用户等级": "user_level",
            "图片URL": "image_url",
        }

        imported_count = 0
        failed_count = 0
        errors = []

        with get_db_session() as session:
            for row_idx, row in enumerate(ws.iter_rows(min_row=3, values_only=True), start=3):
                # 跳过空行
                if not any(row):
                    continue

                data = {}
                for col_idx, value in enumerate(row):
                    if col_idx < len(headers) and headers[col_idx] in header_map:
                        field = header_map[headers[col_idx]]
                        data[field] = value

                # 验证必填字段
                name = data.get("name")
                category = data.get("category")

                if not name or not category:
                    failed_count += 1
                    errors.append({"row": row_idx, "message": "缺少必填字段"})
                    continue

                if category not in ["hook", "sinker", "swivel", "leader", "float", "snap", "other"]:
                    failed_count += 1
                    errors.append({"row": row_idx, "message": f"无效的分类: {category}"})
                    continue

                # 检查名称是否已存在
                existing = session.query(Accessory).filter(Accessory.name == name).first()
                if existing:
                    failed_count += 1
                    errors.append({"row": row_idx, "message": f"配件 '{name}' 已存在"})
                    continue

                # 创建配件
                try:
                    accessory = Accessory(
                        name=str(name),
                        category=str(category),
                        description=str(data.get("description")) if data.get("description") else None,
                        features=str(data.get("features")) if data.get("features") else None,
                        size=str(data.get("size")) if data.get("size") else None,
                        weight=float(data.get("weight")) if data.get("weight") else None,
                        material=str(data.get("material")) if data.get("material") else None,
                        color=str(data.get("color")) if data.get("color") else None,
                        quantity_per_pack=int(data.get("quantity_per_pack")) if data.get("quantity_per_pack") else None,
                        target_species=str(data.get("target_species")) if data.get("target_species") else None,
                        applicable_rigs=str(data.get("applicable_rigs")) if data.get("applicable_rigs") else None,
                        best_conditions=str(data.get("best_conditions")) if data.get("best_conditions") else None,
                        brand=str(data.get("brand")) if data.get("brand") else None,
                        price_min=float(data.get("price_min")) if data.get("price_min") else None,
                        price_max=float(data.get("price_max")) if data.get("price_max") else None,
                        user_level=str(data.get("user_level")) if data.get("user_level") else "beginner",
                        image_url=str(data.get("image_url")) if data.get("image_url") else None,
                    )
                    session.add(accessory)
                    imported_count += 1
                except Exception as e:
                    failed_count += 1
                    errors.append({"row": row_idx, "message": str(e)})

            session.commit()

        logger.info(
            f"配件导入完成: imported={imported_count}, failed={failed_count}, "
            f"user={current_user.username}"
        )

        return AccessoryImportResult(
            success=imported_count > 0,
            message=f"成功导入 {imported_count} 条，失败 {failed_count} 条",
            imported_count=imported_count,
            failed_count=failed_count,
            errors=errors[:10],  # 只返回前10个错误
        )
    except ImportError:
        raise HTTPException(status_code=500, detail="未安装 openpyxl，无法处理 Excel 文件")
    except Exception as e:
        logger.error(f"导入失败: {e}")
        raise HTTPException(status_code=400, detail=f"导入失败: {str(e)}")


@router.get(
    "/accessories/export",
    summary="导出配件数据",
)
async def export_accessories(
    current_user=Depends(get_current_user),
):
    """导出所有配件数据为Excel文件"""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

        with get_db_session() as session:
            accessories = session.query(Accessory).order_by(Accessory.category, Accessory.name).all()

        wb = Workbook()
        ws = wb.active
        ws.title = "配件数据"

        # 样式定义
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin'),
        )

        # 定义列
        columns = [
            ("accessory_id", "ID"),
            ("name", "配件名称"),
            ("category", "分类"),
            ("description", "描述"),
            ("features", "特点"),
            ("size", "规格尺寸"),
            ("weight", "重量(g)"),
            ("material", "材质"),
            ("color", "颜色"),
            ("quantity_per_pack", "每包数量"),
            ("target_species", "目标鱼种"),
            ("applicable_rigs", "适用钓组"),
            ("best_conditions", "最佳条件"),
            ("brand", "品牌"),
            ("price_min", "最低价格"),
            ("price_max", "最高价格"),
            ("user_level", "用户等级"),
            ("image_url", "图片URL"),
        ]

        # 写入表头
        for col_idx, (_, label) in enumerate(columns, 1):
            cell = ws.cell(row=1, column=col_idx, value=label)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center')
            cell.border = border
            ws.column_dimensions[cell.column_letter].width = 15

        # 写入数据
        for row_idx, accessory in enumerate(accessories, 2):
            for col_idx, (field, _) in enumerate(columns, 1):
                value = getattr(accessory, field, None)
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.border = border

        # 保存到内存
        output = BytesIO()
        wb.save(output)
        output.seek(0)

        return StreamingResponse(
            output,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={
                "Content-Disposition": f"attachment; filename=accessories_export.xlsx"
            },
        )
    except ImportError:
        raise HTTPException(status_code=500, detail="未安装 openpyxl，无法生成导出文件")
    except Exception as e:
        logger.error(f"导出失败: {e}")
        raise HTTPException(status_code=500, detail=f"导出失败: {str(e)}")


# ========== Init Data Endpoint ==========


@router.post(
    "/accessories/init",
    response_model=AccessoryInitDataResponse,
    summary="初始化默认配件数据",
)
async def init_accessories(
    current_user=Depends(require_permission("content:create")),
):
    """初始化常用钓鱼配件数据（仅在数据库为空时执行）"""
    with get_db_session() as session:
        # 检查是否已有数据
        existing_count = session.query(Accessory).count()
        if existing_count > 0:
            return AccessoryInitDataResponse(
                created_count=0,
                message=f"数据库已有 {existing_count} 条配件数据，跳过初始化",
            )

        # 创建默认数据
        created_count = 0
        for accessory_data in DEFAULT_ACCESSORIES:
            accessory = Accessory(**accessory_data)
            session.add(accessory)
            created_count += 1

        session.commit()

        logger.info(
            f"配件数据初始化完成: created_count={created_count}, "
            f"user={current_user.username}"
        )

        return AccessoryInitDataResponse(
            created_count=created_count,
            message=f"成功创建 {created_count} 条配件数据",
        )
