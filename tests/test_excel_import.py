#!/usr/bin/env python3
"""
Excel 导入功能测试

测试内容：
1. ExcelImportService - 模板生成、Excel解析、数据导入
2. Excel 导入 API 端点 - 模板下载、预览、执行
3. 审核来源筛选功能
"""

import io
import json
import logging
import sys
from pathlib import Path
from typing import Generator
from unittest.mock import MagicMock, patch

import pytest

# 添加项目路径
sys.path.insert(0, str(Path(__file__).parent.parent))

# 配置日志
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)


# ========== Fixtures ==========


@pytest.fixture
def excel_import_service():
    """创建 ExcelImportService 实例"""
    from apps.api.services.excel_import_service import ExcelImportService
    return ExcelImportService()


@pytest.fixture
def sample_rod_excel_data():
    """示例鱼竿 Excel 数据"""
    return [
        {
            "品牌名称": "达瓦",
            "产品名称": "银狼纪念版",
            "型号": "SL-01",
            "最低价格": 1500,
            "最高价格": 1800,
            "产品描述": "高端路亚竿",
        },
        {
            "品牌名称": "禧玛诺",
            "产品名称": "炎月",
            "型号": "YY-02",
            "最低价格": 800,
            "最高价格": 1200,
            "产品描述": "中端路亚竿",
        },
    ]


# ========== ExcelImportService 单元测试 ==========


class TestExcelImportService:
    """ExcelImportService 单元测试"""

    def test_generate_rod_template(self, excel_import_service):
        """测试生成鱼竿模板"""
        template_bytes = excel_import_service.generate_template("rod")

        assert template_bytes is not None
        assert len(template_bytes) > 0
        # Excel 文件以 PK 开头（ZIP 格式）
        assert template_bytes[:2] == b"PK"

    def test_generate_reel_template(self, excel_import_service):
        """测试生成渔轮模板"""
        template_bytes = excel_import_service.generate_template("reel")

        assert template_bytes is not None
        assert len(template_bytes) > 0
        assert template_bytes[:2] == b"PK"

    def test_generate_line_template(self, excel_import_service):
        """测试生成鱼线模板"""
        template_bytes = excel_import_service.generate_template("line")

        assert template_bytes is not None
        assert len(template_bytes) > 0

    def test_generate_lure_template(self, excel_import_service):
        """测试生成拟饵模板"""
        template_bytes = excel_import_service.generate_template("lure")

        assert template_bytes is not None
        assert len(template_bytes) > 0

    def test_generate_invalid_template(self, excel_import_service):
        """测试生成无效类型模板"""
        with pytest.raises(ValueError) as exc_info:
            excel_import_service.generate_template("invalid_type")

        assert "不支持的装备类型" in str(exc_info.value)

    def test_parse_valid_rod_excel(self, excel_import_service):
        """测试解析有效的鱼竿 Excel - 使用模板格式"""
        import openpyxl

        # 先生成模板
        template_bytes = excel_import_service.generate_template("rod")

        # 加载模板并添加数据
        wb = openpyxl.load_workbook(io.BytesIO(template_bytes))
        ws = wb.active

        # 从第4行开始填入数据（前3行是表头和说明）
        ws.cell(row=4, column=1, value="达瓦")  # brand_name
        ws.cell(row=4, column=2, value="银狼纪念版")  # name
        ws.cell(row=4, column=3, value="SL-01")  # model
        ws.cell(row=4, column=4, value=1500)  # price_min
        ws.cell(row=4, column=5, value=1800)  # price_max

        ws.cell(row=5, column=1, value="禧玛诺")
        ws.cell(row=5, column=2, value="炎月")
        ws.cell(row=5, column=3, value="YY-02")
        ws.cell(row=5, column=4, value=800)
        ws.cell(row=5, column=5, value=1200)

        # 保存到字节流
        excel_bytes = io.BytesIO()
        wb.save(excel_bytes)
        file_content = excel_bytes.getvalue()

        # 解析 Excel
        result = excel_import_service.parse_excel(file_content, "rod")

        assert len(result) == 2
        # 检查第一条数据
        first_row = result[0]
        assert first_row.is_valid is True
        assert first_row.data["brand_name"] == "达瓦"
        assert first_row.data["name"] == "银狼纪念版"

    def test_parse_excel_with_missing_required_fields(self, excel_import_service):
        """测试解析缺少必填字段的 Excel"""
        import openpyxl

        # 生成模板并只填写品牌，缺少产品名称
        template_bytes = excel_import_service.generate_template("rod")
        wb = openpyxl.load_workbook(io.BytesIO(template_bytes))
        ws = wb.active

        # 只填品牌
        ws.cell(row=4, column=1, value="达瓦")

        excel_bytes = io.BytesIO()
        wb.save(excel_bytes)
        file_content = excel_bytes.getvalue()

        result = excel_import_service.parse_excel(file_content, "rod")

        assert len(result) >= 1
        # 第一行应该有错误
        assert result[0].is_valid is False
        assert len(result[0].errors) > 0

    def test_parse_empty_excel(self, excel_import_service):
        """测试解析空 Excel（只有表头无数据）"""
        # 生成模板（没有添加数据）
        template_bytes = excel_import_service.generate_template("rod")

        result = excel_import_service.parse_excel(template_bytes, "rod")

        # 应该返回空列表
        assert len(result) == 0

    def test_validate_row_internal(self, excel_import_service):
        """测试内部验证方法"""
        # 有效数据
        valid_data = {
            "brand_name": "达瓦",
            "name": "银狼纪念版",
            "model": "SL-01",
            "price_min": "1500",
            "price_max": "1800",
        }
        errors = excel_import_service._validate_row(valid_data, "鱼竿")
        assert len(errors) == 0

        # 缺少品牌
        missing_brand = {
            "name": "银狼纪念版",
            "model": "SL-01",
        }
        errors = excel_import_service._validate_row(missing_brand, "鱼竿")
        assert len(errors) > 0
        assert any("品牌" in e for e in errors)

        # 缺少产品名称
        missing_name = {
            "brand_name": "达瓦",
            "model": "SL-01",
        }
        errors = excel_import_service._validate_row(missing_name, "鱼竿")
        assert len(errors) > 0
        assert any("产品名称" in e for e in errors)

    def test_build_extracted_data(self, excel_import_service):
        """测试构建提取数据结构"""
        row_data = {
            "brand_name": "达瓦",
            "name": "银狼纪念版",
            "model": "SL-01",
            "price_min": "1500",
            "price_max": "1800",
            "description": "高端路亚竿",
            "features": "轻量化,高灵敏度",
            "target_fish": "鲈鱼,翘嘴",
            "user_level": "进阶",
            "spec_length": "2.1m",
            "spec_power": "ML",
        }

        result = excel_import_service._build_extracted_data(row_data, "鱼竿")

        assert result["brand_name"] == "达瓦"
        assert result["name"] == "银狼纪念版"
        assert result["price_min"] == 1500.0
        assert result["price_max"] == 1800.0
        assert "轻量化" in result["features"]
        assert "鲈鱼" in result["target_fish"]
        assert result["specs"]["length"] == "2.1m"
        assert result["specs"]["power"] == "ML"

    def test_template_has_correct_structure(self, excel_import_service):
        """测试模板结构正确"""
        import openpyxl

        template_bytes = excel_import_service.generate_template("rod")
        wb = openpyxl.load_workbook(io.BytesIO(template_bytes))
        ws = wb.active

        # 检查第一行是标签
        assert ws.cell(row=1, column=1).value == "品牌名称"
        assert ws.cell(row=1, column=2).value == "产品名称"

        # 检查第三行是字段名
        assert ws.cell(row=3, column=1).value == "brand_name"
        assert ws.cell(row=3, column=2).value == "name"


# ========== Excel 导入 API 端点测试 ==========


class TestExcelImportAPI:
    """Excel 导入 API 端点测试"""

    @pytest.fixture
    def api_client(self):
        """创建 API 测试客户端"""
        from fastapi.testclient import TestClient
        from apps.api.main import app

        return TestClient(app)

    @pytest.fixture
    def auth_token(self, api_client):
        """获取认证 token"""
        import bcrypt
        from apps.api.orm.session import get_db_session, init_db
        from apps.api.orm.repositories.admin_user_repo import AdminUserRepository

        # 初始化数据库
        init_db(create_tables=True)

        # 创建管理员用户
        try:
            with get_db_session() as session:
                repo = AdminUserRepository(session)
                existing = repo.get_by_username("admin")
                if not existing:
                    password_hash = bcrypt.hashpw(
                        "admin123".encode(), bcrypt.gensalt()
                    ).decode()
                    repo.create_admin_user(
                        username="admin",
                        email="admin@test.com",
                        password_hash=password_hash,
                        role="admin",
                    )
        except Exception as e:
            logger.warning(f"创建管理员用户失败: {e}")

        # 登录获取 token
        response = api_client.post(
            "/api/v1/auth/login", json={"username": "admin", "password": "admin123"}
        )

        if response.status_code == 200:
            return response.json()["access_token"]
        return None

    def test_get_import_templates_list(self, api_client, auth_token):
        """测试获取导入模板列表"""
        if not auth_token:
            pytest.skip("无法获取认证 token")

        response = api_client.get(
            "/api/v1/admin/workflow/import/templates",
            headers={"Authorization": f"Bearer {auth_token}"},
        )

        assert response.status_code == 200
        data = response.json()
        assert "templates" in data
        assert len(data["templates"]) == 4  # rod, reel, line, lure

        # 检查每个模板的结构
        for template in data["templates"]:
            assert "equipment_type" in template
            assert "equipment_type_label" in template
            assert "download_url" in template

    def test_download_rod_template(self, api_client, auth_token):
        """测试下载鱼竿模板"""
        if not auth_token:
            pytest.skip("无法获取认证 token")

        response = api_client.get(
            "/api/v1/admin/workflow/import/template/rod",
            headers={"Authorization": f"Bearer {auth_token}"},
        )

        assert response.status_code == 200
        assert (
            response.headers["content-type"]
            == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        content_disposition = response.headers.get("content-disposition", "")
        assert "attachment" in content_disposition
        # RFC 5987 编码格式: filename*=UTF-8''xxx
        assert "filename*=UTF-8''" in content_disposition or "filename=" in content_disposition

        # 验证是有效的 Excel 文件
        content = response.content
        assert content[:2] == b"PK"

    def test_download_invalid_template(self, api_client, auth_token):
        """测试下载无效类型模板"""
        if not auth_token:
            pytest.skip("无法获取认证 token")

        response = api_client.get(
            "/api/v1/admin/workflow/import/template/invalid",
            headers={"Authorization": f"Bearer {auth_token}"},
        )

        assert response.status_code == 400

    def test_preview_import_valid_excel(self, api_client, auth_token):
        """测试预览有效的 Excel 导入"""
        if not auth_token:
            pytest.skip("无法获取认证 token")

        import openpyxl
        from apps.api.services.excel_import_service import get_excel_import_service

        # 使用模板创建测试数据
        service = get_excel_import_service()
        template_bytes = service.generate_template("rod")

        wb = openpyxl.load_workbook(io.BytesIO(template_bytes))
        ws = wb.active

        # 填入数据
        ws.cell(row=4, column=1, value="达瓦")
        ws.cell(row=4, column=2, value="银狼纪念版")
        ws.cell(row=4, column=3, value="SL-01")
        ws.cell(row=4, column=4, value=1500)
        ws.cell(row=4, column=5, value=1800)

        excel_bytes = io.BytesIO()
        wb.save(excel_bytes)
        excel_bytes.seek(0)

        response = api_client.post(
            "/api/v1/admin/workflow/import/preview?equipment_type=rod",
            headers={"Authorization": f"Bearer {auth_token}"},
            files={"file": ("test.xlsx", excel_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )

        assert response.status_code == 200
        data = response.json()
        assert data["success"] is True
        assert data["total_rows"] == 1
        assert data["valid_rows"] == 1
        assert data["invalid_rows"] == 0

    def test_preview_import_invalid_excel(self, api_client, auth_token):
        """测试预览无效的 Excel 导入"""
        if not auth_token:
            pytest.skip("无法获取认证 token")

        import openpyxl
        from apps.api.services.excel_import_service import get_excel_import_service

        # 使用模板但只填写品牌（缺少产品名称）
        service = get_excel_import_service()
        template_bytes = service.generate_template("rod")

        wb = openpyxl.load_workbook(io.BytesIO(template_bytes))
        ws = wb.active
        ws.cell(row=4, column=1, value="达瓦")  # 只有品牌

        excel_bytes = io.BytesIO()
        wb.save(excel_bytes)
        excel_bytes.seek(0)

        response = api_client.post(
            "/api/v1/admin/workflow/import/preview?equipment_type=rod",
            headers={"Authorization": f"Bearer {auth_token}"},
            files={"file": ("test.xlsx", excel_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )

        assert response.status_code == 200
        data = response.json()
        assert data["success"] is True
        assert data["invalid_rows"] >= 1

    def test_execute_import_valid_excel(self, api_client, auth_token):
        """测试执行有效的 Excel 导入"""
        if not auth_token:
            pytest.skip("无法获取认证 token")

        import openpyxl
        from apps.api.services.excel_import_service import get_excel_import_service

        # 使用模板创建测试数据
        service = get_excel_import_service()
        template_bytes = service.generate_template("rod")

        wb = openpyxl.load_workbook(io.BytesIO(template_bytes))
        ws = wb.active

        ws.cell(row=4, column=1, value="测试品牌")
        ws.cell(row=4, column=2, value="测试产品")
        ws.cell(row=4, column=3, value="TEST-001")
        ws.cell(row=4, column=4, value=100)
        ws.cell(row=4, column=5, value=200)

        excel_bytes = io.BytesIO()
        wb.save(excel_bytes)
        excel_bytes.seek(0)

        response = api_client.post(
            "/api/v1/admin/workflow/import/execute?equipment_type=rod",
            headers={"Authorization": f"Bearer {auth_token}"},
            files={"file": ("test.xlsx", excel_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )

        assert response.status_code == 200
        data = response.json()
        assert data["success"] is True
        assert data["imported_count"] >= 1


# ========== 审核来源筛选测试 ==========


class TestReviewSourceFilter:
    """审核来源筛选测试"""

    @pytest.fixture
    def api_client(self):
        """创建 API 测试客户端"""
        from fastapi.testclient import TestClient
        from apps.api.main import app

        return TestClient(app)

    @pytest.fixture
    def auth_token(self, api_client):
        """获取认证 token"""
        import bcrypt
        from apps.api.orm.session import get_db_session, init_db
        from apps.api.orm.repositories.admin_user_repo import AdminUserRepository

        init_db(create_tables=True)

        try:
            with get_db_session() as session:
                repo = AdminUserRepository(session)
                existing = repo.get_by_username("admin")
                if not existing:
                    password_hash = bcrypt.hashpw(
                        "admin123".encode(), bcrypt.gensalt()
                    ).decode()
                    repo.create_admin_user(
                        username="admin",
                        email="admin@test.com",
                        password_hash=password_hash,
                        role="admin",
                    )
        except Exception as e:
            logger.warning(f"创建管理员用户失败: {e}")

        response = api_client.post(
            "/api/v1/auth/login", json={"username": "admin", "password": "admin123"}
        )

        if response.status_code == 200:
            return response.json()["access_token"]
        return None

    def test_get_review_tasks_all(self, api_client, auth_token):
        """测试获取所有审核任务"""
        if not auth_token:
            pytest.skip("无法获取认证 token")

        response = api_client.get(
            "/api/v1/admin/workflow/review/tasks",
            headers={"Authorization": f"Bearer {auth_token}"},
        )

        assert response.status_code == 200
        data = response.json()
        assert "items" in data
        assert "total" in data
        assert "page" in data

    def test_get_review_tasks_filter_by_source_type(self, api_client, auth_token):
        """测试按来源类型筛选审核任务"""
        if not auth_token:
            pytest.skip("无法获取认证 token")

        # 测试筛选 excel_import 来源
        response = api_client.get(
            "/api/v1/admin/workflow/review/tasks?source_type=excel_import",
            headers={"Authorization": f"Bearer {auth_token}"},
        )

        assert response.status_code == 200
        data = response.json()
        assert "items" in data

        # 如果有结果，验证所有项都是 excel_import 来源
        for item in data["items"]:
            assert item.get("source_type") == "excel_import"

    def test_get_review_tasks_filter_by_ecommerce(self, api_client, auth_token):
        """测试按电商来源筛选"""
        if not auth_token:
            pytest.skip("无法获取认证 token")

        response = api_client.get(
            "/api/v1/admin/workflow/review/tasks?source_type=ecommerce",
            headers={"Authorization": f"Bearer {auth_token}"},
        )

        assert response.status_code == 200
        data = response.json()
        assert "items" in data

    def test_get_review_tasks_filter_by_status_and_source(self, api_client, auth_token):
        """测试同时按状态和来源筛选"""
        if not auth_token:
            pytest.skip("无法获取认证 token")

        response = api_client.get(
            "/api/v1/admin/workflow/review/tasks?status=pending&source_type=excel_import",
            headers={"Authorization": f"Bearer {auth_token}"},
        )

        assert response.status_code == 200
        data = response.json()
        assert "items" in data

        # 验证筛选结果
        for item in data["items"]:
            assert item.get("status") == "pending"
            assert item.get("source_type") == "excel_import"


# ========== 集成测试 ==========


class TestExcelImportIntegration:
    """Excel 导入集成测试"""

    @pytest.fixture
    def api_client(self):
        """创建 API 测试客户端"""
        from fastapi.testclient import TestClient
        from apps.api.main import app

        return TestClient(app)

    @pytest.fixture
    def auth_token(self, api_client):
        """获取认证 token"""
        import bcrypt
        from apps.api.orm.session import get_db_session, init_db
        from apps.api.orm.repositories.admin_user_repo import AdminUserRepository

        init_db(create_tables=True)

        try:
            with get_db_session() as session:
                repo = AdminUserRepository(session)
                existing = repo.get_by_username("admin")
                if not existing:
                    password_hash = bcrypt.hashpw(
                        "admin123".encode(), bcrypt.gensalt()
                    ).decode()
                    repo.create_admin_user(
                        username="admin",
                        email="admin@test.com",
                        password_hash=password_hash,
                        role="admin",
                    )
        except Exception as e:
            logger.warning(f"创建管理员用户失败: {e}")

        response = api_client.post(
            "/api/v1/auth/login", json={"username": "admin", "password": "admin123"}
        )

        if response.status_code == 200:
            return response.json()["access_token"]
        return None

    def test_full_import_workflow(self, api_client, auth_token):
        """测试完整的导入工作流：下载模板 -> 预览 -> 执行 -> 筛选"""
        if not auth_token:
            pytest.skip("无法获取认证 token")

        import openpyxl

        # 1. 下载模板
        response = api_client.get(
            "/api/v1/admin/workflow/import/template/rod",
            headers={"Authorization": f"Bearer {auth_token}"},
        )
        assert response.status_code == 200, "下载模板失败"
        template_bytes = response.content

        # 2. 使用模板创建带数据的 Excel
        wb = openpyxl.load_workbook(io.BytesIO(template_bytes))
        ws = wb.active

        # 添加测试数据（从第4行开始）
        test_data = [
            ("集成测试品牌A", "集成测试产品1", "INT-001", 500, 800, "集成测试描述1"),
            ("集成测试品牌B", "集成测试产品2", "INT-002", 600, 900, "集成测试描述2"),
        ]

        for row_idx, (brand, name, model, price_min, price_max, desc) in enumerate(test_data, 4):
            ws.cell(row=row_idx, column=1, value=brand)
            ws.cell(row=row_idx, column=2, value=name)
            ws.cell(row=row_idx, column=3, value=model)
            ws.cell(row=row_idx, column=4, value=price_min)
            ws.cell(row=row_idx, column=5, value=price_max)
            ws.cell(row=row_idx, column=6, value=desc)

        excel_bytes = io.BytesIO()
        wb.save(excel_bytes)

        # 3. 预览导入
        excel_bytes.seek(0)
        response = api_client.post(
            "/api/v1/admin/workflow/import/preview?equipment_type=rod",
            headers={"Authorization": f"Bearer {auth_token}"},
            files={"file": ("test.xlsx", excel_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert response.status_code == 200, "预览导入失败"
        preview_data = response.json()
        assert preview_data["valid_rows"] == 2, f"预览数据不正确: {preview_data}"

        # 4. 执行导入
        excel_bytes.seek(0)
        response = api_client.post(
            "/api/v1/admin/workflow/import/execute?equipment_type=rod",
            headers={"Authorization": f"Bearer {auth_token}"},
            files={"file": ("test.xlsx", excel_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert response.status_code == 200, "执行导入失败"
        import_result = response.json()
        assert import_result["success"] is True, f"导入不成功: {import_result}"
        assert import_result["imported_count"] == 2, f"导入数量不正确: {import_result}"

        # 5. 验证可以在审核列表中筛选到
        response = api_client.get(
            "/api/v1/admin/workflow/review/tasks?source_type=excel_import",
            headers={"Authorization": f"Bearer {auth_token}"},
        )
        assert response.status_code == 200, "获取审核列表失败"
        review_data = response.json()

        # 检查是否有 excel_import 来源的数据
        assert review_data["total"] >= 2, "审核列表中应该有导入的数据"

        logger.info("完整导入工作流测试通过！")


def main():
    """运行测试"""
    import subprocess
    import sys

    result = subprocess.run(
        [sys.executable, "-m", "pytest", __file__, "-v", "--tb=short"],
        capture_output=False,
    )
    return result.returncode


if __name__ == "__main__":
    exit(main())
